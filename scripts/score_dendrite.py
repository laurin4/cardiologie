#!/usr/bin/env python3
"""Score cardiology_smoke predictions against Dendrite postop gold labels."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from configs.config import OUTPUTS_DIR, PREDICTIONS_DIR, RAW_DATA_DIR
from src.evaluation.dendrite_score import (
    PAIR_COLUMNS,
    discover_dendrite_paths,
    format_score_report,
    run_dendrite_score,
)


def _write_pairs_excel(df: pd.DataFrame, out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows

    wb = Workbook()
    overview = wb.active
    overview.title = "all"
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c_idx, value in enumerate(row, start=1):
            cell = overview.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.font = header_font
            else:
                cell.alignment = wrap

    widths = {
        "fall": 14,
        "patient_id": 12,
        "field": 22,
        "dendrite_column": 28,
        "gold_raw": 36,
        "gold": 14,
        "pred_raw": 18,
        "pred": 14,
        "match": 10,
        "scored": 10,
        "exclude_reason": 28,
        "source_report": 22,
        "source_columns": 36,
        "evidence_quotes": 55,
        "reasoning": 40,
    }
    for col_idx, name in enumerate(df.columns, start=1):
        overview.column_dimensions[get_column_letter(col_idx)].width = widths.get(
            str(name), 16
        )
    overview.freeze_panes = "A2"
    overview.auto_filter.ref = overview.dimensions

    if "field" in df.columns:
        for field, sub in df.groupby("field", sort=True):
            title = str(field)[:28]
            ws = wb.create_sheet(title=title)
            for r_idx, row in enumerate(
                dataframe_to_rows(sub, index=False, header=True), start=1
            ):
                for c_idx, value in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)
                    if r_idx == 1:
                        cell.font = header_font
                    else:
                        cell.alignment = wrap
            for col_idx, name in enumerate(sub.columns, start=1):
                ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(
                    str(name), 16
                )
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Score pipeline results vs Dendrite postop gold (FallNummer join)."
    )
    parser.add_argument(
        "--predictions",
        default=None,
        help="Pipeline results CSV (default: outputs/extractions_dendrite/cardiology_smoke_results.csv).",
    )
    parser.add_argument(
        "--dendrite",
        default=None,
        help="Dendrite gold Excel/CSV (default: discover Dendrite* under data/raw/).",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Where to write score JSON/CSV/XLSX (default: outputs/evaluation/).",
    )
    parser.add_argument(
        "--max-patients",
        type=int,
        default=25,
        help=(
            "Number of patients in the review grid (default: 25). "
            "Same FallNummern for every diagnosis → equal row counts. "
            "Use 0 for all overlapping patients."
        ),
    )
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument(
        "--complete-only",
        action="store_true",
        help=(
            "Legacy: export only rows with both gold and scorable pred "
            "(unequal counts per field). Default is the equal patient grid."
        ),
    )
    args = parser.parse_args()

    pred_path = (
        Path(args.predictions)
        if args.predictions
        else Path("outputs/extractions_dendrite/cardiology_smoke_results.csv")
    )
    if not pred_path.exists():
        alt = PREDICTIONS_DIR / "cardiology_smoke_results.csv"
        if alt.exists():
            pred_path = alt
        else:
            raise SystemExit(f"Predictions not found: {pred_path}")

    if args.dendrite:
        dend_path = Path(args.dendrite)
    else:
        found = discover_dendrite_paths(RAW_DATA_DIR)
        if not found:
            raise SystemExit(
                f"No Dendrite* file under {RAW_DATA_DIR}. Pass --dendrite explicitly."
            )
        dend_path = found[0]
        if len(found) > 1:
            print(f"Multiple Dendrite files; using {dend_path.name}")

    if not dend_path.exists():
        raise SystemExit(f"Dendrite file not found: {dend_path}")

    max_patients = None if args.max_patients == 0 else args.max_patients
    result = run_dendrite_score(
        pred_path,
        dend_path,
        max_patients=max_patients,
        seed=args.seed,
        complete_only=args.complete_only,
    )
    report = format_score_report(result)
    print(report)
    n_export = result.get("n_patients_in_export", {})
    print(
        f"Export rows per field (equal patient grid unless --complete-only): {n_export} "
        f"| unique falls={result.get('n_unique_falls_export')} "
        f"(max-patients={args.max_patients}, seed={args.seed}, "
        f"complete_only={args.complete_only})"
    )
    counts = list(n_export.values()) if n_export else []
    if counts and len(set(counts)) > 1 and not args.complete_only:
        print("WARNING: unequal field counts — unexpected for patient-grid mode.")
    if counts and all(c == 0 for c in counts):
        raise SystemExit("Export has 0 rows.")
    empty_ev = 0
    for rows in (result.get("pairs") or {}).values():
        empty_ev += sum(1 for r in rows if not str(r.get("evidence_quotes") or "").strip())
    if empty_ev:
        print(
            f"NOTE: {empty_ev} export rows still have empty evidence_quotes "
            "(no LLM quotes and no keyword hits in source text)."
        )

    out_dir = Path(args.output_dir) if args.output_dir else OUTPUTS_DIR / "evaluation"
    out_dir.mkdir(parents=True, exist_ok=True)
    summary = {
        "predictions": str(pred_path),
        "dendrite": str(dend_path),
        "n_aligned_fall": result["n_aligned_fall"],
        "n_patients_in_export": result.get("n_patients_in_export"),
        "complete_only": result.get("complete_only", False),
        "n_unique_falls_export": result.get("n_unique_falls_export"),
        "max_patients": args.max_patients,
        "seed": args.seed,
        "per_field": result["per_field"],
    }
    json_path = out_dir / "dendrite_score.json"
    json_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    pair_rows = []
    for field, rows in (result.get("pairs") or {}).items():
        pair_rows.extend(rows)
    df = pd.DataFrame(pair_rows, columns=PAIR_COLUMNS) if pair_rows else pd.DataFrame(
        columns=PAIR_COLUMNS
    )
    pairs_path = out_dir / "dendrite_score_pairs.csv"
    df.to_csv(
        pairs_path,
        index=False,
        sep=";",
        encoding="utf-8-sig",
        quoting=1,  # csv.QUOTE_ALL — prevents column shift in Excel
    )
    xlsx_path = out_dir / "dendrite_score_pairs.xlsx"
    _write_pairs_excel(df, xlsx_path)

    print(f"Wrote {json_path}")
    print(f"Wrote {pairs_path} (semicolon, all fields quoted)")
    print(f"Wrote {xlsx_path}  ← use this file for review")
    if result["n_aligned_fall"] == 0:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
