"""
Rodney review export: idiot-proof long-format rows with report type, columns, snippets.

Sample up to N cases per clinical variable from Dendrite-overlap predictions.
"""

from __future__ import annotations

import csv
import random
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from src.evaluation.review_sheet import CLINICAL_COLUMNS, _clean, _quotes_cell
from src.preprocessing.verlegung_loader import provenance_for_text_source

RODNEY_COLUMNS = [
    "variable",
    "patient_id",
    "fall_nummers",
    "verlegung_fallnr",
    "prediction",
    "source_report",
    "source_columns",
    "evidence_quotes",
    "reasoning",
    "status",
    "correct",
    "notes",
]


def _text_source_for_variable(var_name: str) -> str:
    try:
        from configs.tasks import load_task

        task = load_task("cardiology_smoke")
        for var in task.variables or ():
            if var.name == var_name:
                return getattr(var, "text_source", "report") or "report"
    except Exception:
        pass
    defaults = {
        "pacemaker": "verlegung",
        "atrial_fibrillation": "both",
        "cerebrovascular_event": "both",
        "reoperation_required": "both",
        "reoperation_context": "both",
        "multi_system_failure": "verlegung",
        "rethoracotomy": "verlegung",
        "rethoracotomy_context": "verlegung",
        "liver_cirrhosis": "austritt",
    }
    return defaults.get(var_name, "report")


def build_long_rows(result_rows: list[dict], variables: Optional[Sequence[str]] = None) -> list[dict]:
    vars_ = list(variables) if variables is not None else list(CLINICAL_COLUMNS)
    out: list[dict] = []
    for row in result_rows:
        patient_id = (
            _clean(row.get("patient_id"))
            or _clean(row.get("report_id"))
            or _clean(row.get("source_row_id"))
        )
        for var in vars_:
            src_report = _clean(row.get(f"{var}_source_report"))
            src_cols = _clean(row.get(f"{var}_source_columns"))
            if not src_report:
                src_report, src_cols = provenance_for_text_source(_text_source_for_variable(var))
            quotes = _quotes_cell(row.get(f"{var}_evidence_quotes") or row.get("evidence_quotes"))
            reasoning = _clean(row.get(f"{var}_reasoning") or "")
            if not reasoning and var == "reoperation_required":
                reasoning = _clean(row.get("reasoning")).replace("\n", " ")
            out.append(
                {
                    "variable": var,
                    "patient_id": patient_id,
                    "fall_nummers": _clean(row.get("fall_nummers")),
                    "verlegung_fallnr": _clean(row.get("verlegung_fallnr")),
                    "prediction": _clean(row.get(var)),
                    "source_report": src_report,
                    "source_columns": src_cols,
                    "evidence_quotes": quotes,
                    "reasoning": reasoning.replace("\n", " "),
                    "status": _clean(row.get("status")),
                    "correct": "",
                    "notes": "",
                }
            )
    return out


def sample_patients(
    result_rows: list[dict],
    *,
    n_patients: int = 25,
    seed: int = 42,
) -> list[dict]:
    """
    Sample up to *n_patients* unique patients, then keep ALL their variable rows.

    Ensures every variable has the same patient set / same row count.
    """
    if n_patients <= 0 or len(result_rows) <= n_patients:
        return list(result_rows)

    rng = random.Random(seed)
    # Stable unique key per patient row
    keyed: List[tuple[str, dict]] = []
    seen: set[str] = set()
    for row in result_rows:
        key = (
            _clean(row.get("verlegung_fallnr"))
            or _clean(row.get("patient_id"))
            or _clean(row.get("report_id"))
            or _clean(row.get("fall_nummers"))
        )
        if not key or key in seen:
            continue
        seen.add(key)
        keyed.append((key, row))

    if len(keyed) <= n_patients:
        return [r for _, r in keyed]

    picked_keys = {k for k, _ in rng.sample(keyed, n_patients)}
    # Preserve original order among selected
    out = []
    for key, row in keyed:
        if key in picked_keys:
            out.append(row)
    return out


def sample_per_variable(
    long_rows: list[dict],
    *,
    n_per_var: int = 25,
    seed: int = 42,
    variables: Optional[Sequence[str]] = None,
) -> list[dict]:
    """Stratify lightly by prediction value, then sample up to n_per_var per variable."""
    vars_ = list(variables) if variables is not None else list(CLINICAL_COLUMNS)
    rng = random.Random(seed)
    sampled: list[dict] = []
    for var in vars_:
        pool = [r for r in long_rows if r.get("variable") == var]
        if not pool:
            continue
        by_pred: Dict[str, List[dict]] = {}
        for r in pool:
            by_pred.setdefault(r.get("prediction") or "", []).append(r)
        buckets = [list(v) for v in by_pred.values()]
        for b in buckets:
            rng.shuffle(b)
        picked: list[dict] = []
        # Round-robin across prediction buckets for diversity.
        while len(picked) < n_per_var and any(buckets):
            progressed = False
            for b in buckets:
                if not b:
                    continue
                picked.append(b.pop())
                progressed = True
                if len(picked) >= n_per_var:
                    break
            if not progressed:
                break
        sampled.extend(picked)
    return sampled


def filter_dendrite_overlap(
    result_rows: list[dict],
    dendrite_falls: set[str],
) -> list[dict]:
    from src.preprocessing.report_identity import normalize_str

    kept = []
    for row in result_rows:
        falls = []
        raw = _clean(row.get("fall_nummers"))
        if raw:
            falls.extend(normalize_str(x) for x in raw.replace("|", ";").split(";") if x.strip())
        vfall = normalize_str(row.get("verlegung_fallnr") or "")
        if vfall:
            falls.append(vfall)
        rid = normalize_str(row.get("report_id") or row.get("patient_id") or "")
        if rid:
            falls.append(rid)
        if any(f and f in dendrite_falls for f in falls):
            kept.append(row)
    return kept


def write_rodney_csv(rows: list[dict], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=RODNEY_COLUMNS,
            delimiter=";",
            extrasaction="ignore",
            quoting=csv.QUOTE_MINIMAL,
        )
        writer.writeheader()
        writer.writerows(rows)


def write_rodney_excel(rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    # One sheet per variable + overview
    by_var: Dict[str, List[dict]] = {}
    for r in rows:
        by_var.setdefault(str(r.get("variable") or "unknown"), []).append(r)

    overview = wb.active
    overview.title = "overview"
    overview["A1"] = "variable"
    overview["B1"] = "n_rows"
    overview["A1"].font = Font(bold=True)
    overview["B1"].font = Font(bold=True)
    for i, (var, items) in enumerate(sorted(by_var.items()), start=2):
        overview.cell(row=i, column=1, value=var)
        overview.cell(row=i, column=2, value=len(items))

    wrap = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True)
    widths = {
        "variable": 22,
        "patient_id": 14,
        "fall_nummers": 18,
        "verlegung_fallnr": 14,
        "prediction": 22,
        "source_report": 22,
        "source_columns": 36,
        "evidence_quotes": 50,
        "reasoning": 40,
        "status": 12,
        "correct": 12,
        "notes": 30,
    }

    for var, items in sorted(by_var.items()):
        title = var[:28] if len(var) <= 28 else var[:25] + "..."
        ws = wb.create_sheet(title=title)
        for col_idx, name in enumerate(RODNEY_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=name)
            cell.font = header_font
        for row_idx, row in enumerate(items, start=2):
            for col_idx, name in enumerate(RODNEY_COLUMNS, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(name, ""))
                cell.alignment = wrap
        for col_idx, name in enumerate(RODNEY_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(name, 16)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
